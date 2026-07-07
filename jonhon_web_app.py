#!/usr/bin/env python3
"""Web API for drawing upload, OCR, material recommendation, and template export."""

from __future__ import annotations

import json
import os
import shutil
import threading
import time
import uuid
from concurrent.futures import ThreadPoolExecutor
from pathlib import Path
from typing import Any

import requests
from fastapi import BackgroundTasks, FastAPI, File, Form, HTTPException, UploadFile
from fastapi.responses import FileResponse
from fastapi.staticfiles import StaticFiles
from openpyxl import load_workbook

import jonhon_material_pipeline_no_fallback as pipeline


APP_DIR = Path(__file__).resolve().parent
STATIC_DIR = APP_DIR / "web_frontend"
RUNS_DIR = APP_DIR / "治理输出" / "web_runs"
KNOWLEDGE_UPLOADS_DIR = APP_DIR / "治理输出" / "web_knowledge_files"
KNOWLEDGE_MANIFEST = KNOWLEDGE_UPLOADS_DIR / "manifest.json"
DEFAULT_KNOWLEDGE_ID = "__default__"
REQUIRED_KNOWLEDGE_SHEETS = [
    "导线电缆知识",
    "导线标准线径",
    "防护辅材知识",
    "扎带卡扣知识",
    "连接器BOM映射",
]
DEFAULT_OCR_URL = os.getenv("JONHON_OCR_URL", "http://39.96.162.45:18130/ocr")
OCR_TIMEOUT_SECONDS = int(os.getenv("JONHON_OCR_TIMEOUT", "300"))
DEFAULT_LLM_MODEL = os.getenv("JONHON_WEB_LLM_MODEL", "qwen3.6-max-preview")

app = FastAPI(title="JONHON Material Pipeline", version="0.1.0")
app.mount("/static", StaticFiles(directory=STATIC_DIR), name="static")

executor = ThreadPoolExecutor(max_workers=int(os.getenv("JONHON_WEB_WORKERS", "2")))
jobs: dict[str, dict[str, Any]] = {}
jobs_lock = threading.Lock()
KNOWLEDGE_UPLOADS_DIR.mkdir(parents=True, exist_ok=True)


@app.get("/")
def index() -> FileResponse:
    return FileResponse(STATIC_DIR / "index.html")


@app.get("/api/health")
def health() -> dict[str, Any]:
    return {"ok": True, "ocr_url": DEFAULT_OCR_URL, "default_knowledge": pipeline.DEFAULT_KNOWLEDGE_XLSX.name}


@app.get("/api/knowledge-files")
def list_knowledge_files() -> dict[str, Any]:
    return {"items": get_knowledge_files()}


@app.post("/api/knowledge-files")
async def upload_knowledge_file(file: UploadFile = File(...)) -> dict[str, Any]:
    if not file.filename:
        raise HTTPException(status_code=400, detail="缺少上传文件名")
    suffix = Path(file.filename).suffix.lower()
    if suffix != ".xlsx":
        raise HTTPException(status_code=400, detail="治理文件必须是 .xlsx")

    file_id = time.strftime("%Y%m%d_%H%M%S_") + uuid.uuid4().hex[:8] + suffix
    target = KNOWLEDGE_UPLOADS_DIR / file_id
    with target.open("wb") as out:
        shutil.copyfileobj(file.file, out)

    try:
        validate_knowledge_workbook(target)
    except Exception as exc:
        target.unlink(missing_ok=True)
        raise HTTPException(status_code=400, detail=str(exc)) from exc

    manifest = load_knowledge_manifest()
    manifest[file_id] = {
        "id": file_id,
        "name": sanitize_filename(file.filename),
        "uploaded_at": time.time(),
        "size": target.stat().st_size,
    }
    save_knowledge_manifest(manifest)
    return {"item": knowledge_file_record(file_id, target, manifest[file_id])}


@app.post("/api/jobs")
async def create_job(
    background_tasks: BackgroundTasks,
    file: UploadFile = File(...),
    use_llm_extract: bool = Form(False),
    knowledge_file_id: str = Form(DEFAULT_KNOWLEDGE_ID),
) -> dict[str, Any]:
    if not file.filename:
        raise HTTPException(status_code=400, detail="缺少上传文件名")
    suffix = Path(file.filename).suffix.lower()
    if suffix not in {".pdf", ".docx"}:
        raise HTTPException(status_code=400, detail="目前只支持 PDF 图纸或 DOCX 技术要求文件")

    job_id = time.strftime("%Y%m%d_%H%M%S_") + uuid.uuid4().hex[:8]
    job_dir = RUNS_DIR / job_id
    job_dir.mkdir(parents=True, exist_ok=True)
    input_file = job_dir / sanitize_filename(file.filename)
    with input_file.open("wb") as out:
        shutil.copyfileobj(file.file, out)
    knowledge_path, knowledge_label = resolve_knowledge_file(knowledge_file_id)

    init_job(
        job_id,
        {
            "job_id": job_id,
            "status": "queued",
            "stage": "等待处理",
            "created_at": time.time(),
            "updated_at": time.time(),
            "input_file": str(input_file),
            "job_dir": str(job_dir),
            "input_type": suffix.lstrip("."),
            "use_llm_extract": use_llm_extract,
            "knowledge_file_id": knowledge_file_id,
            "knowledge_file": str(knowledge_path),
            "knowledge_file_name": knowledge_label,
            "files": {},
            "summary": {},
            "error": "",
        },
    )
    background_tasks.add_task(submit_job, job_id, input_file, job_dir, use_llm_extract, knowledge_path)
    return {"job_id": job_id, "status_url": f"/api/jobs/{job_id}"}


@app.get("/api/jobs/{job_id}")
def get_job(job_id: str) -> dict[str, Any]:
    job = read_job(job_id)
    if not job:
        raise HTTPException(status_code=404, detail="任务不存在")
    return public_job(job)


@app.get("/api/jobs/{job_id}/download/{file_key}")
def download_file(job_id: str, file_key: str) -> FileResponse:
    job = read_job(job_id)
    if not job:
        raise HTTPException(status_code=404, detail="任务不存在")
    files = job.get("files") or {}
    path = files.get(file_key)
    if not path:
        raise HTTPException(status_code=404, detail="文件不存在")
    file_path = Path(path)
    if not file_path.exists():
        raise HTTPException(status_code=404, detail="文件已不存在")
    return FileResponse(file_path, filename=file_path.name)


def submit_job(job_id: str, input_file: Path, job_dir: Path, use_llm_extract: bool, knowledge_xlsx: Path) -> None:
    executor.submit(run_job, job_id, input_file, job_dir, use_llm_extract, knowledge_xlsx)


def run_job(job_id: str, input_file: Path, job_dir: Path, use_llm_extract: bool, knowledge_xlsx: Path) -> None:
    try:
        source_type = input_file.suffix.lower().lstrip(".")
        ocr_text_path = job_dir / "ocr_normalized.txt"
        ocr_tables_path = job_dir / "ocr_tables.json"
        files: dict[str, str] = {}

        if source_type == "docx":
            update_job(job_id, status="running", stage="解析 DOCX 文本和表格")
            ocr_text, tables = pipeline.load_docx_inputs(input_file)
            docx_tables_path = job_dir / "docx_tables.json"
            pipeline.write_json(docx_tables_path, tables)
            files["docx_tables"] = str(docx_tables_path)
        else:
            update_job(job_id, status="running", stage="调用 OCR 服务")
            ocr_json = job_dir / "ocr_result.json"
            call_ocr_service(input_file, ocr_json)
            files["ocr_json"] = str(ocr_json)

            update_job(job_id, stage="解析 OCR 表格")
            blocks = load_ocr_blocks_compat(ocr_json)
            ocr_text = pipeline.blocks_to_ocr_text(blocks)
            tables = pipeline.extract_html_tables(blocks)

        ocr_text_path.write_text(ocr_text, encoding="utf-8")
        pipeline.write_json(ocr_tables_path, tables)
        files["ocr_text"] = str(ocr_text_path)

        update_job(job_id, stage="抽取图纸需求")
        req = build_requirements(ocr_text, tables, use_llm_extract, source_type=source_type)

        update_job(job_id, stage="治理表匹配和物料推荐")
        rec = pipeline.build_recommendations(req, knowledge_xlsx)
        requirements_path = job_dir / "requirements.json"
        recommendations_path = job_dir / "recommendations.json"
        pipeline.write_json(requirements_path, pipeline.asdict(req))
        pipeline.write_json(recommendations_path, rec)

        update_job(job_id, stage="回填模板")
        output_xlsx = job_dir / "线缆组件产品接线表_自动生成.xlsx"
        pipeline.make_output_workbook(pipeline.DEFAULT_TEMPLATE_XLS, output_xlsx, req, rec)

        summary = build_summary(req, rec)
        files.update(
            {
                "output": str(output_xlsx),
                "requirements": str(requirements_path),
                "recommendations": str(recommendations_path),
            }
        )
        update_job(
            job_id,
            status="done",
            stage="完成",
            files=files,
            summary=summary,
        )
    except Exception as exc:
        update_job(job_id, status="failed", stage="失败", error=str(exc))


def call_ocr_service(input_pdf: Path, output_json: Path) -> None:
    with input_pdf.open("rb") as fh:
        response = requests.post(DEFAULT_OCR_URL, files={"file": fh}, timeout=OCR_TIMEOUT_SECONDS)
    response.raise_for_status()
    try:
        payload = response.json()
    except Exception as exc:
        raise RuntimeError(f"OCR 服务未返回 JSON: {response.text[:500]}") from exc
    output_json.write_text(json.dumps(payload, ensure_ascii=False, indent=2), encoding="utf-8")


def load_ocr_blocks_compat(ocr_json: Path) -> list[dict[str, Any]]:
    data = json.loads(ocr_json.read_text(encoding="utf-8"))
    try:
        return pipeline.load_ocr_blocks(ocr_json)
    except Exception:
        pass
    if isinstance(data, dict) and isinstance(data.get("pages"), list):
        blocks: list[dict[str, Any]] = []
        order = 0
        for page in data["pages"]:
            page_index = page.get("page_index", 0)
            content = page.get("content") or []
            if not isinstance(content, list):
                continue
            for item in content:
                if not isinstance(item, dict):
                    continue
                text = pipeline.norm_text(item.get("text"))
                label = pipeline.norm_text(item.get("type")) or "text"
                if not text:
                    continue
                blocks.append(
                    {
                        "block_id": f"p{page_index}_{order}",
                        "block_order": order,
                        "block_label": label,
                        "block_content": text,
                        "block_bbox": item.get("bbox") or [],
                        "page_index": page_index,
                    }
                )
                order += 1
        if blocks:
            return blocks
    raise ValueError(f"无法在 {ocr_json} 中找到可解析的 OCR 内容")


def build_requirements(
    ocr_text: str,
    tables: list[dict[str, Any]],
    use_llm_extract: bool,
    source_type: str = "ocr",
) -> pipeline.DrawingRequirements:
    extractor = pipeline.LLMExtractor(model=DEFAULT_LLM_MODEL) if use_llm_extract else None
    return pipeline.build_requirements_from_text(
        ocr_text,
        tables,
        use_llm_extract=use_llm_extract,
        extractor=extractor,
        source_type=source_type,
    )


def build_summary(req: pipeline.DrawingRequirements, rec: dict[str, Any]) -> dict[str, Any]:
    selected = rec.get("selected") or {}
    materials = pipeline.build_connector_sheet_materials(rec)
    wire = selected.get("wire") or {}
    protection = selected.get("protection") or {}
    return {
        "connection_rows": len(req.connection_rows),
        "drawing_materials": len(req.materials_from_drawing),
        "sheet_materials": len(materials),
        "wire": wire.get("物料编号") or "",
        "protection": protection.get("物料编号") or "",
        "connectors": len(selected.get("connectors") or []),
        "connector_bom_items": len(selected.get("connector_bom_items") or []),
        "tie_clips": len(selected.get("tie_clips") or []),
    }


def sanitize_filename(filename: str) -> str:
    name = Path(filename).name.replace("/", "_").replace("\\", "_")
    return name or "drawing.pdf"


def get_knowledge_files() -> list[dict[str, Any]]:
    manifest = load_knowledge_manifest()
    items = [
        knowledge_file_record(
            DEFAULT_KNOWLEDGE_ID,
            pipeline.DEFAULT_KNOWLEDGE_XLSX,
            {
                "name": pipeline.DEFAULT_KNOWLEDGE_XLSX.name,
                "uploaded_at": None,
                "size": pipeline.DEFAULT_KNOWLEDGE_XLSX.stat().st_size if pipeline.DEFAULT_KNOWLEDGE_XLSX.exists() else 0,
            },
            is_default=True,
        )
    ]
    for path in sorted(KNOWLEDGE_UPLOADS_DIR.glob("*.xlsx"), key=lambda p: p.stat().st_mtime, reverse=True):
        meta = manifest.get(path.name, {})
        items.append(knowledge_file_record(path.name, path, meta))
    return items


def knowledge_file_record(
    file_id: str,
    path: Path,
    meta: dict[str, Any] | None = None,
    *,
    is_default: bool = False,
) -> dict[str, Any]:
    meta = meta or {}
    stat = path.stat() if path.exists() else None
    return {
        "id": file_id,
        "name": meta.get("name") or path.name,
        "is_default": is_default,
        "size": meta.get("size") or (stat.st_size if stat else 0),
        "updated_at": meta.get("uploaded_at") or (stat.st_mtime if stat else None),
    }


def resolve_knowledge_file(file_id: str) -> tuple[Path, str]:
    if not file_id or file_id == DEFAULT_KNOWLEDGE_ID:
        validate_knowledge_workbook(pipeline.DEFAULT_KNOWLEDGE_XLSX)
        return pipeline.DEFAULT_KNOWLEDGE_XLSX, pipeline.DEFAULT_KNOWLEDGE_XLSX.name
    if Path(file_id).name != file_id:
        raise HTTPException(status_code=400, detail="治理文件 ID 非法")
    candidate = KNOWLEDGE_UPLOADS_DIR / file_id
    if not candidate.exists():
        raise HTTPException(status_code=404, detail="治理文件不存在")
    validate_knowledge_workbook(candidate)
    manifest = load_knowledge_manifest()
    return candidate, manifest.get(file_id, {}).get("name") or candidate.name


def validate_knowledge_workbook(path: Path) -> None:
    if not path.exists():
        raise ValueError(f"治理文件不存在: {path}")
    try:
        wb = load_workbook(path, read_only=True, data_only=True)
    except Exception as exc:
        raise ValueError("无法读取治理文件，请确认它是有效的 .xlsx 文件") from exc
    try:
        missing = [sheet for sheet in REQUIRED_KNOWLEDGE_SHEETS if sheet not in wb.sheetnames]
        if missing:
            raise ValueError("治理文件缺少必要 sheet: " + "、".join(missing))
    finally:
        wb.close()


def load_knowledge_manifest() -> dict[str, Any]:
    if not KNOWLEDGE_MANIFEST.exists():
        return {}
    try:
        data = json.loads(KNOWLEDGE_MANIFEST.read_text(encoding="utf-8"))
    except Exception:
        return {}
    return data if isinstance(data, dict) else {}


def save_knowledge_manifest(manifest: dict[str, Any]) -> None:
    KNOWLEDGE_MANIFEST.write_text(json.dumps(manifest, ensure_ascii=False, indent=2), encoding="utf-8")


def init_job(job_id: str, job: dict[str, Any]) -> None:
    with jobs_lock:
        jobs[job_id] = job


def read_job(job_id: str) -> dict[str, Any] | None:
    with jobs_lock:
        job = jobs.get(job_id)
        return dict(job) if job else None


def update_job(job_id: str, **updates: Any) -> None:
    with jobs_lock:
        job = jobs[job_id]
        job.update(updates)
        job["updated_at"] = time.time()


def public_job(job: dict[str, Any]) -> dict[str, Any]:
    result = dict(job)
    result.pop("input_pdf", None)
    result.pop("input_file", None)
    result.pop("job_dir", None)
    result.pop("knowledge_file", None)
    files = result.get("files") or {}
    result["downloads"] = {key: f"/api/jobs/{job['job_id']}/download/{key}" for key in files}
    return result


if __name__ == "__main__":
    import uvicorn

    uvicorn.run(
        "jonhon_web_app:app",
        host=os.getenv("JONHON_WEB_HOST", "0.0.0.0"),
        port=int(os.getenv("JONHON_WEB_PORT", "8050")),
        reload=False,
    )
